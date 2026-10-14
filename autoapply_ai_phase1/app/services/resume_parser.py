from io import BytesIO
import fitz
from docx import Document


def extract_text(filename: str, content: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "pdf":
        with fitz.open(stream=content, filetype="pdf") as pdf:
            return "\n".join(page.get_text() for page in pdf).strip()
    if suffix == "docx":
        document = Document(BytesIO(content))
        chunks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(chunks).strip()
    raise ValueError("Unsupported resume format. Upload a PDF or DOCX.")
